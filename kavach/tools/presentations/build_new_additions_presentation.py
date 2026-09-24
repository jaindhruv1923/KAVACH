from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOCS_ROOT = PROJECT_ROOT.parents[1] / "docs" / "presentations"
SOURCE = DOCS_ROOT / "KAVACH_Presentation_Updated.pptx"
OUTPUT = DOCS_ROOT / "new_additions.pptx"

NAVY = RGBColor(11, 18, 32)
PANEL = RGBColor(20, 31, 52)
PANEL2 = RGBColor(18, 48, 57)
BLUE = RGBColor(82, 156, 255)
CYAN = RGBColor(79, 214, 207)
WHITE = RGBColor(241, 246, 252)
MUTED = RGBColor(166, 183, 204)
GREEN = RGBColor(71, 207, 137)
AMBER = RGBColor(246, 181, 74)
RED = RGBColor(240, 100, 100)


def blank_slide(prs):
    layout = next((item for item in prs.slide_layouts if not item.placeholders), prs.slide_layouts[-1])
    slide = prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_text(slide, "KAVACH  /  NEW ADDITIONS", 0.55, 0.28, 8.0, 0.24, 9, CYAN, True)
    add_text(slide, f"{len(prs.slides):02d}", 12.25, 7.2, 0.4, 0.22, 9, MUTED, True, PP_ALIGN.RIGHT)
    return slide


def add_text(slide, text, left, top, width, height, size=16, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.07)
    frame.margin_right = Inches(0.07)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def panel(slide, left, top, width, height, fill=PANEL, line=RGBColor(54, 81, 119)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def heading(slide, title, subtitle):
    add_text(slide, title, 0.55, 0.72, 12.0, 0.52, 25, WHITE, True)
    add_text(slide, subtitle, 0.55, 1.32, 11.8, 0.32, 11, MUTED)


def bullets(slide, values, left, top, width, height, size=14, color=MUTED):
    add_text(slide, "\n".join("•  " + value for value in values), left, top, width, height, size, color)


prs = Presentation(str(SOURCE))

# Slide: release summary
slide = blank_slide(prs)
heading(slide, "What shipped today", "KAVACH moved from a capable prototype to a multi-entry security workspace.")
items = [
    ("01", "Security Command Center", "Live metrics, Gemini status, activity feed, heartbeat, responsive animated cards."),
    ("02", "GitHub Intelligence", "Public repository tree retrieval, source scanning, chunking, and Qdrant indexing."),
    ("03", "Live Code Review", "Standalone PII and credential review with policy risk and remediation guidance."),
    ("04", "Voice-first requests", "Continuous browser recognition with interim transcript updates and clear error states."),
    ("05", "Safer numeric policy", "Context-free numbers are not automatically treated as sensitive identifiers."),
]
for i, (num, title, desc) in enumerate(items):
    y = 1.9 + i * 0.9
    panel(slide, 0.65, y, 12.0, 0.68, PANEL if i % 2 == 0 else PANEL2, CYAN if i == 1 else RGBColor(54, 81, 119))
    add_text(slide, num, 0.92, y + 0.16, 0.42, 0.25, 12, CYAN, True)
    add_text(slide, title, 1.52, y + 0.12, 3.0, 0.27, 14, WHITE, True)
    add_text(slide, desc, 4.7, y + 0.12, 7.45, 0.35, 11, MUTED)

# Slide: before/after comparison
slide = blank_slide(prs)
heading(slide, "Before vs after: the product evolution", "The security core remains stable; the product surface is now broader, clearer, and easier to operate.")
panel(slide, 0.65, 1.85, 5.75, 4.75, RGBColor(27, 35, 51), RGBColor(87, 98, 119))
panel(slide, 6.55, 1.85, 6.1, 4.75, PANEL2, CYAN)
add_text(slide, "EARLIER BASELINE", 0.95, 2.15, 3.0, 0.25, 10, AMBER, True)
add_text(slide, "Governed workflow prototype", 0.95, 2.5, 4.6, 0.38, 19, WHITE, True)
bullets(slide, [
    "Local repository ingestion and RAG",
    "Agent planning and orchestration",
    "PII, secret, and policy checks",
    "Impact analysis and grounded generation",
    "Validation and workflow trace",
], 1.0, 3.15, 4.75, 2.6, 14)
add_text(slide, "CURRENT RELEASE", 6.88, 2.15, 3.0, 0.25, 10, CYAN, True)
add_text(slide, "Security command center", 6.88, 2.5, 4.6, 0.38, 19, WHITE, True)
bullets(slide, [
    "Multiple entry points: workflow, GitHub, code review",
    "Live operational metrics and activity feedback",
    "Persistent Gemini setup with safe status visibility",
    "Voice requests with live transcript capture",
    "Animated, responsive dashboard for real users",
], 6.93, 3.15, 5.15, 2.6, 14, WHITE)
add_text(slide, "Same governance principle. More usable product.", 1.0, 6.82, 11.5, 0.32, 14, WHITE, True, PP_ALIGN.CENTER)

# Slide: current architecture
slide = blank_slide(prs)
heading(slide, "The current KAVACH flow", "New interfaces feed the same security and policy engines instead of creating parallel logic.")
steps = [
    ("INPUT", "Developer request\nVoice or text", BLUE),
    ("CONNECT", "Local path or\npublic GitHub", CYAN),
    ("SCAN", "PII + secrets\n+ policy", AMBER),
    ("GROUND", "Qdrant evidence\n+ impact ranking", BLUE),
    ("GENERATE", "Gemini or\nstandalone review", CYAN),
    ("AUDIT", "Validation, trace\n+ decision", GREEN),
]
for i, (label, desc, color) in enumerate(steps):
    x = 0.55 + i * 2.08
    panel(slide, x, 2.25, 1.7, 1.75, PANEL, color)
    add_text(slide, label, x + 0.15, 2.53, 1.4, 0.25, 10, color, True, PP_ALIGN.CENTER)
    add_text(slide, desc, x + 0.12, 3.0, 1.46, 0.65, 13, WHITE, True, PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        add_text(slide, "→", x + 1.77, 2.9, 0.3, 0.4, 18, MUTED, True, PP_ALIGN.CENTER)
add_text(slide, "The new GitHub and code-review paths strengthen the front door; the governed workflow remains the control plane.", 1.0, 4.65, 11.4, 0.45, 15, WHITE, True, PP_ALIGN.CENTER)
panel(slide, 1.0, 5.45, 11.1, 0.78, RGBColor(24, 34, 52), RGBColor(54, 81, 119))
add_text(slide, "Design decision", 1.28, 5.67, 1.5, 0.24, 11, CYAN, True)
add_text(slide, "Reuse one detection and policy layer everywhere so a finding is treated consistently across the product.", 2.9, 5.62, 8.7, 0.3, 13, MUTED)

# Slide: PR Guardian
slide = blank_slide(prs)
heading(slide, "Next-level proposal: KAVACH PR Guardian", "The clearest path from a local workspace to a production-relevant engineering control.")
panel(slide, 0.62, 1.85, 4.0, 4.85, PANEL2, CYAN)
add_text(slide, "THE PROBLEM", 0.95, 2.18, 2.0, 0.25, 10, CYAN, True)
add_text(slide, "AI-generated changes can reach a pull request before anyone understands their data risk, blast radius, or validation status.", 0.95, 2.62, 3.2, 1.1, 19, WHITE, True)
add_text(slide, "The missing control point is before merge.", 0.95, 4.45, 3.1, 0.5, 14, GREEN, True)
add_text(slide, "KAVACH can become that control point.", 0.95, 5.45, 3.15, 0.45, 14, WHITE, True)
panel(slide, 4.95, 1.85, 7.75, 4.85, PANEL, RGBColor(54, 81, 119))
add_text(slide, "PROPOSED PR FLOW", 5.3, 2.18, 2.4, 0.25, 10, AMBER, True)
flow = [
    "1. GitHub App receives pull-request webhook",
    "2. KAVACH scans the diff for secrets, PII, injection, and dangerous commands",
    "3. Repository RAG retrieves surrounding evidence",
    "4. Impact engine ranks affected components",
    "5. Policy engine returns Allow, Review, Redact, or Block",
    "6. KAVACH posts an explainable check to the pull request",
    "7. Human approval is required for high-impact changes",
]
bullets(slide, flow, 5.3, 2.66, 6.8, 3.35, 14, WHITE)
add_text(slide, "Outcome: safer merges without removing developer velocity.", 5.3, 6.2, 6.8, 0.3, 13, GREEN, True)

# Slide: roadmap
slide = blank_slide(prs)
heading(slide, "Roadmap to a production-grade platform", "The recommended sequence turns today's working surfaces into durable organizational controls.")
roadmap = [
    ("PHASE A", "Trustworthy state", "PostgreSQL persistence\nAuthentication + RBAC\nRedacted audit retention", BLUE),
    ("PHASE B", "Change governance", "GitHub App / OAuth\nPR webhooks\nDiff-based scanning\nHuman approvals", CYAN),
    ("PHASE C", "Safe delivery", "Proposed patch view\nIsolated test execution\nDependency + static analysis\nBranch/export after approval", AMBER),
    ("PHASE D", "Enterprise visibility", "Security score trends\nPDF/JSON reports\nPolicy packs\nDocker + observability", GREEN),
]
for i, (phase, title, desc, color) in enumerate(roadmap):
    x = 0.62 + (i % 2) * 6.18
    y = 1.95 + (i // 2) * 2.35
    panel(slide, x, y, 5.65, 1.85, PANEL if i != 1 else PANEL2, color)
    add_text(slide, phase, x + 0.28, y + 0.22, 1.3, 0.22, 10, color, True)
    add_text(slide, title, x + 0.28, y + 0.52, 4.8, 0.3, 18, WHITE, True)
    add_text(slide, desc, x + 0.28, y + 0.96, 4.95, 0.65, 12, MUTED if i != 1 else WHITE)
add_text(slide, "Recommended flagship milestone: KAVACH PR Guardian", 0.8, 6.86, 11.5, 0.35, 15, WHITE, True, PP_ALIGN.CENTER)

# Slide: value and scope
slide = blank_slide(prs)
heading(slide, "Why this becomes a real problem-solving product", "KAVACH is strongest when positioned as a governance layer for AI-assisted software engineering.")
values = [
    ("For regulated teams", "Reduce the chance that customer identifiers, credentials, or sensitive repository context reach an unsafe downstream action."),
    ("For developers", "Keep the speed of AI assistance while making evidence, risk, impact, and validation visible in one place."),
    ("For reviewers", "Replace opaque AI output with explainable decisions, approval points, and a durable audit history."),
    ("For engineering leadership", "Turn AI coding governance into a measurable control with reports, trends, and policy packs."),
]
for i, (title, desc) in enumerate(values):
    y = 1.9 + i * 1.05
    panel(slide, 0.78, y, 11.95, 0.78, PANEL2 if i == 1 else PANEL, CYAN if i == 1 else RGBColor(54, 81, 119))
    add_text(slide, title, 1.08, y + 0.2, 2.65, 0.28, 14, WHITE, True)
    add_text(slide, desc, 3.85, y + 0.17, 8.25, 0.4, 12, MUTED if i != 1 else WHITE)
add_text(slide, "Honest scope", 1.0, 6.45, 1.4, 0.25, 10, AMBER, True)
add_text(slide, "Today's additions are shipped and testable. PR Guardian, persistence, authentication, isolated execution, and enterprise deployment are the next roadmap, not current claims.", 2.55, 6.37, 9.8, 0.42, 12, MUTED)

prs.save(str(OUTPUT))
print(OUTPUT)
