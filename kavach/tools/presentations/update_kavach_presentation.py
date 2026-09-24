from copy import deepcopy
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SOURCE = PROJECT_ROOT.parents[1] / "docs" / "presentations" / "KAVACH_Presentation.pptx"
OUTPUT = SOURCE.with_name("KAVACH_Presentation_Updated.pptx")

NAVY = RGBColor(11, 18, 32)
PANEL = RGBColor(20, 31, 52)
BLUE = RGBColor(82, 156, 255)
CYAN = RGBColor(79, 214, 207)
WHITE = RGBColor(241, 246, 252)
MUTED = RGBColor(166, 183, 204)
GREEN = RGBColor(71, 207, 137)
AMBER = RGBColor(246, 181, 74)
RED = RGBColor(240, 100, 100)


def all_text_shapes(slide):
    return [shape for shape in slide.shapes if hasattr(shape, "text_frame")]


def replace_text(slide, replacements):
    for shape in all_text_shapes(slide):
        original = shape.text
        updated = original
        for old, new in replacements.items():
            updated = updated.replace(old, new)
        if updated != original:
            shape.text = updated


def add_text(slide, text, left, top, width, height, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_panel(slide, left, top, width, height, fill=PANEL, line=RGBColor(48, 75, 112)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_new_slide(prs):
    blank_layout = next((layout for layout in prs.slide_layouts if not layout.placeholders), prs.slide_layouts[-1])
    slide = prs.slides.add_slide(blank_layout)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY

    add_text(slide, "KAVACH  /  SHIPPED PRODUCT UPDATE", 0.55, 0.32, 8.8, 0.28, 10, CYAN, True)
    add_text(slide, "From governed prototype to a usable security workspace", 0.55, 0.72, 12.2, 0.55, 25, WHITE, True)
    add_text(slide, "What changed after the original project presentation", 0.55, 1.32, 11.5, 0.3, 12, MUTED)

    add_panel(slide, 0.55, 1.9, 5.75, 4.65, RGBColor(24, 34, 52), RGBColor(74, 92, 122))
    add_panel(slide, 6.55, 1.9, 6.2, 4.65, RGBColor(18, 48, 57), CYAN)
    add_text(slide, "BEFORE", 0.88, 2.18, 2.0, 0.3, 11, AMBER, True)
    add_text(slide, "A strong integrated prototype", 0.88, 2.52, 4.7, 0.4, 19, WHITE, True)
    before = [
        "Single workflow dashboard for request analysis",
        "Local repository ingestion and RAG evidence",
        "PII and secret detection before generation",
        "Risk-aware policy decisions and audit trace",
        "Gemini generation with local .env support",
    ]
    add_text(slide, "\n".join("•  " + item for item in before), 0.92, 3.16, 4.95, 2.6, 14, MUTED)
    add_text(slide, "FOUNDATION", 0.88, 5.88, 2.0, 0.25, 9, BLUE, True)

    add_text(slide, "NOW", 6.88, 2.18, 2.0, 0.3, 11, CYAN, True)
    add_text(slide, "A complete security command center", 6.88, 2.52, 5.2, 0.4, 19, WHITE, True)
    after = [
        "Live animated command-center dashboard with run metrics",
        "Public GitHub repository ingestion and evidence indexing",
        "Standalone live code review for PII, secrets, and remediation",
        "Persistent Gemini configuration via backend/.env",
        "Reliable browser voice input with live transcript capture",
        "Responsive UI with live activity states and workflow feedback",
    ]
    add_text(slide, "\n".join("•  " + item for item in after), 6.92, 3.16, 5.35, 2.85, 14, WHITE)
    add_text(slide, "SHIPPED", 6.88, 5.88, 2.0, 0.25, 9, GREEN, True)

    add_text(slide, "The core principle stayed the same: every AI-assisted change remains evidence-grounded, policy-checked, explainable, and auditable.", 0.75, 6.86, 11.8, 0.35, 13, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "KAVACH  •  Product evolution", 0.55, 7.25, 4.0, 0.2, 8, MUTED)
    add_text(slide, "24", 12.3, 7.2, 0.35, 0.25, 9, MUTED, True, PP_ALIGN.RIGHT)
    return slide


prs = Presentation(str(SOURCE))

# Keep the original design but correct the top-level claims that no longer match the shipped build.
replace_text(prs.slides[0], {
    "Multilingual Security-Governed": "Security-Governed",
    "Agentic AI DevOps Platform": "Agentic AI DevOps Workspace",
})
replace_text(prs.slides[4], {
    "including code-mixed and multilingual requests": "with context-aware security checks",
    "Multilingual PII and secrets detection": "Context-aware PII and secrets detection",
})
replace_text(prs.slides[5], {
    "Multilingual PII detection": "Context-aware PII detection",
    "Hindi/Marathi content": "English security context and Indian identifiers",
})
replace_text(prs.slides[6], {
    "Multilingual PII and secrets detection with severity and explanations": "Context-aware PII and secrets detection with severity and explanations",
})
replace_text(prs.slides[8], {
    "Works across English, Hindi and Marathi text": "Uses explicit English security context and identifier formats",
})
replace_text(prs.slides[14], {
    "Multilingual test corpus": "English security test corpus",
})
replace_text(prs.slides[15], {
    "A multilingual test corpus to check retrieval": "A security test corpus to check retrieval",
})
replace_text(prs.slides[18], {
    "Multilingual PII detection with severity, confidence and explainability": "Context-aware PII detection with severity, confidence and explainability",
})
replace_text(prs.slides[19], {
    "Broader multilingual datasets and red-team evaluation": "Broader datasets and red-team evaluation",
    "Aadhaar-like: 12 digits AND the word Aadhaar (or its Hindi form) nearby": "Aadhaar-like: 12 digits with explicit identity context nearby",
})
replace_text(prs.slides[21], {
    "Phases 4-9, from security hardening through finalization, will be walked through in detail at the next review.": "Phases 4-9 are now integrated, demonstrated, and extended with the shipped command-center workspace.",
    "Delivered and demoable": "Delivered and demoable",
})

add_new_slide(prs)
prs.save(str(OUTPUT))
print(OUTPUT)
